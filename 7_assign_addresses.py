import os
import zipfile
from xml.etree import ElementTree as ET
import pandas as pd
from math import radians, cos, sin, sqrt, atan2
from openpyxl import load_workbook

# Directories
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BASE_DIR, "data")
OUTPUT_DIR = os.path.join(BASE_DIR, "output")

# Auto-detect the KMZ file in the data folder
kmz_files = [f for f in os.listdir(DATA_DIR) if f.lower().endswith(".kmz")]
if len(kmz_files) != 1:
    raise FileNotFoundError(f"Expected 1 KMZ file in {DATA_DIR}, found {len(kmz_files)}")
KMZ_PATH = os.path.join(DATA_DIR, kmz_files[0])
print(f"Using KMZ file: {kmz_files[0]}")

# Static paths for input/output
CONNECTIONS_PATH = os.path.join(DATA_DIR, "Connections_Table.xlsx")
PROCESSED_XLSX_PATH = os.path.join(OUTPUT_DIR, "Combined_Formatted_Output_processed.xlsx")
FINAL_OUTPUT_PATH = os.path.join(OUTPUT_DIR, "Combined_Formatted_Output_with_Addresses.xlsx")

# Haversine distance
def haversine(lat1, lon1, lat2, lon2):
    R = 6371000
    phi1, phi2 = radians(lat1), radians(lat2)
    d_phi = radians(lat2 - lat1)
    d_lambda = radians(lon2 - lon1)
    a = sin(d_phi / 2) ** 2 + cos(phi1) * cos(phi2) * sin(d_lambda / 2) ** 2
    return 2 * R * atan2(sqrt(a), sqrt(1 - a))

# Safe coordinate parser
def safe_extract_coordinates(text):
    try:
        coords = text.strip().replace('\n', ' ').split()
        lon, lat = map(float, coords[0].split(",")[:2])
        return lat, lon
    except:
        return None, None

# Extract addresses from folder <name> and first placemark <coordinates>
def extract_addresses_from_kmz(kmz_path):
    addresses = []
    with zipfile.ZipFile(kmz_path, 'r') as kmz:
        kml_file = next((f for f in kmz.namelist() if f.endswith(".kml")), None)
        if not kml_file:
            raise FileNotFoundError("No KML file found inside the KMZ.")
        kml_content = kmz.read(kml_file)
        root = ET.fromstring(kml_content)

    ns = {"kml": "http://www.opengis.net/kml/2.2"}

    for folder in root.findall(".//kml:Folder", ns):
        folder_name_elem = folder.find("kml:name", ns)
        if folder_name_elem is None:
            continue
        folder_name = folder_name_elem.text.strip()

        if "MI" not in folder_name or "," not in folder_name:
            continue

        placemark = folder.find(".//kml:Placemark", ns)
        coords_elem = placemark.find(".//kml:coordinates", ns) if placemark is not None else None
        if coords_elem is None:
            continue

        lat, lon = safe_extract_coordinates(coords_elem.text)
        if lat is not None and lon is not None:
            addresses.append({
                "Address": folder_name,
                "Latitude": lat,
                "Longitude": lon
            })

    print(f"Extracted {len(addresses)} address placemarks from KMZ.")
    return addresses

def main():
    try:
        # Load coordinates from Connections Table
        print("Loading connection coordinates...")
        df_conn = pd.read_excel(CONNECTIONS_PATH)
        location_coords = {
            row["Location"]: (row["Latitude"], row["Longitude"])
            for _, row in df_conn.iterrows()
        }
        print(f"Loaded {len(location_coords)} locations.")

        # Extract addresses
        addresses = extract_addresses_from_kmz(KMZ_PATH)

        # Load processed workbook
        print("Loading output workbook...")
        wb = load_workbook(PROCESSED_XLSX_PATH)

        updated_count = 0
        for sheet_name in wb.sheetnames:
            if sheet_name not in location_coords:
                continue
            lat1, lon1 = location_coords[sheet_name]
            best_addr = None
            best_dist = float("inf")
            for addr in addresses:
                dist = haversine(lat1, lon1, addr["Latitude"], addr["Longitude"])
                if dist < best_dist:
                    best_addr = addr["Address"]
                    best_dist = dist
            if best_addr:
                ws = wb[sheet_name]
                ws["B5"] = best_addr
                updated_count += 1

        # Save
        wb.save(FINAL_OUTPUT_PATH)
        print(f"✅ Updated {updated_count} sheets.")
        print(f"📁 Output saved to: {FINAL_OUTPUT_PATH}")

    except Exception as e:
        print(f"❌ Error: {e}")

if __name__ == "__main__":
    main()

